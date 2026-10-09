"""Build the GuardianAI CoCo Quest 2026 showcase deck (12 slides, 16:9, with speaker notes).

Run:  python3 ppt/build_deck.py   (writes ppt/coco_quest.pptx)
"""
import os
import shutil
import tempfile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOT = os.path.join(ROOT, "screenshots")
ARCH = os.path.join(ROOT, "docs", "architecture.png")
OUT = os.path.join(ROOT, "ppt", "coco_quest.pptx")

# ---- Theme: light navy + Snowflake blue -------------------------------------
BG = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x0B, 0x25, 0x45)
NAVY2 = RGBColor(0x1B, 0x3A, 0x63)
SF_BLUE = RGBColor(0x29, 0xB5, 0xE8)
SF_MID = RGBColor(0x11, 0x56, 0x7F)
SF_LIGHT = RGBColor(0xE3, 0xF5, 0xFC)
TEXT = RGBColor(0x1F, 0x2D, 0x3D)
MUTED = RGBColor(0x5B, 0x6B, 0x7F)
SOFT = RGBColor(0xC9, 0xD6, 0xE6)
BORDER = RGBColor(0xD6, 0xE2, 0xEE)
AMBER = RGBColor(0xF5, 0xA6, 0x23)
AMBER_LIGHT = RGBColor(0xFD, 0xF1, 0xDC)
AMBER_DARK = RGBColor(0x8A, 0x5A, 0x00)
RED = RGBColor(0xE5, 0x48, 0x4D)
GREEN = RGBColor(0x2B, 0xB6, 0x73)
FONT = "Calibri"
MONO = "Consolas"
TOTAL = 12

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---- Helpers ----------------------------------------------------------------
def bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, italic=False, spacing=1.1, shape=None):
    """content: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) runs."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        for r in (p if isinstance(p, list) else [(p, {})]):
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = t
            f = run.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return shape


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, gap=8, dot=SF_BLUE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(gap)
        r0 = para.add_run()
        r0.text = "\u25A0  "
        r0.font.size = Pt(size - 4)
        r0.font.color.rgb = dot
        r0.font.name = FONT
        r = para.add_run()
        r.text = it
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb


def arrow_line(slide, x1, y1, x2, y2, color=SF_MID, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def header(slide, title, kicker, n):
    bg(slide)
    box(slide, 0.6, 0.55, 0.09, 0.62, fill=SF_BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.82, 0.38, 10, 0.35, kicker.upper(), size=12, color=SF_MID, bold=True)
    text(slide, 0.82, 0.62, 11.8, 0.7, title, size=30, color=NAVY, bold=True)
    box(slide, 0, 7.12, 13.333, 0.38, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.6, 7.15, 9, 0.32, [[("GuardianAI", {"bold": True, "color": WHITE}),
                                       ("   |   CoCo Quest 2026   |   Biswajit Jena", {"color": SOFT})]],
         size=11, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 11.7, 7.15, 1.05, 0.32, f"{n} / {TOTAL}", size=11, color=SF_BLUE, bold=True,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    p.line.color.rgb = BORDER
    p.line.width = Pt(1)
    return p


def pill(slide, x, y, w, h, label, fill=SF_LIGHT, color=SF_MID, size=11):
    s = box(slide, x, y, w, h, fill=fill, line=None, radius=0.5)
    text(slide, 0, 0, 0, 0, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=s)
    return s


def badge(slide, x, y, d, label, fill=NAVY, size=14):
    c = box(slide, x, y, d, d, fill=fill, line=None, shape=MSO_SHAPE.OVAL)
    text(slide, 0, 0, 0, 0, label, size=size, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)


def cell(slide, tbl, i, j, value, fill, color, size=13, bold=False, align=PP_ALIGN.LEFT):
    c = tbl.cell(i, j)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    text(slide, 0, 0, 0, 0, value, size=size, color=color, bold=bold, align=align, shape=c)
    c.vertical_anchor = MSO_ANCHOR.MIDDLE


# =============================================================================
# 1. Title
# =============================================================================
s = prs.slides.add_slide(BLANK)
bg(s)
box(s, 0, 0, 0.35, 7.5, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
box(s, 0.35, 0, 0.08, 7.5, fill=SF_BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
pill(s, 1.0, 1.0, 5.0, 0.42, "CoCo Quest 2026  |  Theme 1: Agentic Data Quality Guardian", size=12)
text(s, 0.95, 1.65, 7.6, 1.1, [[("Guardian", {"color": NAVY}), ("AI", {"color": SF_BLUE})]], size=66, bold=True)
text(s, 1.0, 2.85, 7.4, 0.6, "Autonomous Data Steward for Snowflake", size=28, color=NAVY2, bold=True)
text(s, 1.0, 3.6, 7.0, 1.3,
     "Turns messy, growing data into trusted data: autonomously, continuously and safely. "
     "A live data trust score recovers the moment a human approves a fix.",
     size=17, color=MUTED, spacing=1.2)
text(s, 1.0, 5.55, 7, 0.4, "Biswajit Jena", size=20, color=NAVY, bold=True)
text(s, 1.0, 5.95, 7, 0.4, "Solo entry", size=14, color=MUTED)

box(s, 8.75, 1.35, 3.85, 4.75, fill=WHITE, line=BORDER, radius=0.06)
box(s, 8.75, 1.35, 3.85, 0.12, fill=SF_BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, 8.9, 1.7, 3.55, 0.4, "DATA TRUST SCORE", size=13, color=SF_MID, bold=True, align=PP_ALIGN.CENTER)
text(s, 8.85, 2.15, 3.65, 1.2, [[("62", {"color": RED}), (" \u2192 ", {"color": MUTED, "size": 36}),
                                 ("~100", {"color": GREEN})]], size=56, bold=True, align=PP_ALIGN.CENTER)
text(s, 8.9, 3.35, 3.55, 0.4, "Overall, across 5 related tables", size=13, color=MUTED, align=PP_ALIGN.CENTER)
for i, (k, v) in enumerate([("9", "agents"), ("27", "quality checks"), ("5", "related tables")]):
    yy = 4.0 + i * 0.62
    text(s, 9.1, yy, 1.1, 0.55, k, size=26, color=NAVY, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.3, yy, 2.2, 0.55, v, size=15, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Hi, I'm Biswajit Jena, and this is my solo entry for CoCo Quest 2026, Theme 1: Agentic Data Quality "
         "Guardian. GuardianAI is an autonomous data steward that runs entirely inside Snowflake. It detects, "
         "explains, fixes and validates data quality issues across five related tables. It only stops to ask a "
         "human before applying a risky fix. The number to remember is on the right: the overall data trust "
         "score goes from 62 to about 100 once fixes are approved. In the next few minutes I'll show how it "
         "works, why it's agentic, and how it stays safe.")

# =============================================================================
# 2. Problem
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Bad data silently breaks business decisions", "The problem", 2)
for i, (t, d) in enumerate([
    ("Duplicate customers", "Inflate customer counts and every report built on them."),
    ("Orphan orders & payments", "Corrupt revenue attribution across related tables."),
    ("Invalid emails", "Waste marketing spend on customers you can't reach."),
    ("Negative prices", "Distort margin and product profitability."),
]):
    x = 0.6 + i * 3.08
    box(s, x, 1.65, 2.88, 2.35)
    box(s, x, 1.65, 2.88, 0.1, fill=RED, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.2, 1.95, 2.5, 0.75, t, size=19, color=NAVY, bold=True)
    text(s, x + 0.2, 2.75, 2.5, 1.2, d, size=14, color=MUTED, spacing=1.15)

box(s, 0.6, 4.35, 12.13, 2.45, fill=WHITE)
text(s, 0.9, 4.55, 11.5, 0.5, "Traditional tools detect issues, then stop.", size=22, color=NAVY, bold=True)
steps = ["Investigate root cause", "Judge business impact", "Write the fix", "Prove it worked"]
for i, st in enumerate(steps):
    x = 0.9 + i * 2.95
    pill(s, x, 5.25, 2.65, 0.55, st, fill=AMBER_LIGHT, color=AMBER_DARK, size=14)
    if i < len(steps) - 1:
        text(s, x + 2.62, 5.2, 0.35, 0.6, "\u203A", size=26, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.9, 6.05, 11.5, 0.6, "...are all left to a human. That doesn't scale, and it doesn't run continuously.",
     size=16, color=MUTED, italic=True)
notes(s, "Bad data rarely fails loudly. Duplicate customers inflate reports. Orphan orders and payments corrupt "
         "revenue attribution. Invalid emails waste marketing spend, and negative prices distort margins. "
         "Most data quality tools detect these issues and stop there. A person still has to find the root "
         "cause, judge the business impact, write the fix and prove it worked. That manual loop doesn't scale "
         "as data grows, and it doesn't run continuously. That's the gap GuardianAI closes.")

# =============================================================================
# 3. Solution: the loop
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "An agentic loop that owns the full data-quality lifecycle", "The solution", 3)
loop = ["Detect", "Score", "Rebuild Plan", "Explain", "Impact", "Approve", "Remediate", "Validate", "Govern"]
cw, step = 1.41, 1.34
for i, st in enumerate(loop):
    human = st == "Approve"
    c = box(s, 0.6 + i * step, 1.75, cw, 0.95,
            fill=AMBER if human else (NAVY if i % 2 == 0 else SF_MID), line=None,
            shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    text(s, 0, 0, 0, 0, st, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
text(s, 0.6 + 5 * step - 0.2, 2.78, cw + 0.4, 0.4, "Human gate", size=12, color=AMBER_DARK,
     bold=True, align=PP_ALIGN.CENTER)

text(s, 0.6, 3.3, 12, 0.5, "Why it's agentic, not just automation", size=20, color=NAVY, bold=True)
for i, (t, d) in enumerate([
    ("Distinct roles", "Each agent has one job and reasons over the previous agent's output."),
    ("Self-verifying", "The Validation Agent independently re-checks the Remediation Agent's work."),
    ("Self-healing", "The pipeline rebuilds its own fix plan when new issues appear."),
    ("Knows when to stop", "It decides its next step, including when to pause and ask a human."),
]):
    x = 0.6 + i * 3.08
    box(s, x, 3.9, 2.88, 2.9)
    c = box(s, x + 0.22, 4.12, 0.55, 0.55, fill=SF_LIGHT, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 0, 0, 0, 0, str(i + 1), size=16, color=SF_MID, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
    text(s, x + 0.2, 4.8, 2.5, 0.5, t, size=18, color=NAVY, bold=True)
    text(s, x + 0.2, 5.3, 2.5, 1.4, d, size=14, color=MUTED, spacing=1.15)
notes(s, "GuardianAI is a chain of specialised agents, orchestrated by native Snowflake Tasks, that runs the "
         "whole loop: detect, score, rebuild the plan, explain, assess impact, approve, remediate, validate "
         "and govern. The amber step is the only place a human is needed. What makes it agentic rather than a "
         "script: each agent has a distinct role and reasons over the previous agent's output. The Validation "
         "Agent independently re-checks the remediation. The pipeline rebuilds its own fix plan when new "
         "issues appear. And the system decides its own next step, including when to stop and ask a human.")

# =============================================================================
# 4. Architecture
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Architecture: 100% native Snowflake", "How it fits together", 4)
picture(s, ARCH, 4.02, 1.52, h=5.35)
for i, (t, d) in enumerate([
    ("Ingest", "Snowsight upload or landing stage; Streams capture new rows (CDC)."),
    ("Data layer", "CUSTOMERS, ORDERS, PRODUCTS, PAYMENTS, INVENTORY, with cross-table checks."),
    ("Agent pipeline", "Task DAG: Observer \u2192 Score \u2192 Rebuild Plan \u2192 Diagnosis + Impact \u2192 Exec Summary."),
    ("Human gate", "Approve \u2192 Remediate \u2192 Validate \u2192 Govern."),
    ("Presentation", "Streamlit in Snowflake, 5 pages."),
]):
    y = 1.55 + i * 1.08
    box(s, 0.6, y, 3.25, 0.98, fill=AMBER_LIGHT if t == "Human gate" else WHITE)
    text(s, 0.75, y + 0.06, 3.0, 0.35, t, size=14, color=NAVY, bold=True)
    text(s, 0.75, y + 0.38, 3.0, 0.6, d, size=11, color=MUTED, spacing=1.05)
notes(s, "Everything runs natively in Snowflake, with no external services. Data arrives through a Snowsight "
         "upload or a landing stage, and Streams capture the new rows. Five related retail tables form the "
         "data layer. A Task DAG runs the agent pipeline: Observer, Health Scorer, Plan Builder, then the "
         "Cortex-powered Diagnosis, Impact and Executive Summary agents. After that comes the human approval "
         "gate, then Remediate, Validate and Govern. A five-page Streamlit-in-Snowflake app is the window "
         "into all of it.")

# =============================================================================
# 5. The 9 agents
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Nine specialised agents", "The agents", 5)
agents = [
    ("Data Observer", "Deterministic SQL", "27 checks across 5 tables (incl. cross-table) \u2192 DQ_ISSUES"),
    ("Health Scorer", "SQL (volume-aware)", "Severity \u00d7 % rows affected \u2192 per-table + overall score"),
    ("Plan Builder", "SQL", "Self-healing: rebuilds the fix plan for current issues"),
    ("Diagnosis", "Snowflake Cortex", "Plain-English root cause per issue"),
    ("Business Impact", "Snowflake Cortex", "Executive-language business impact"),
    ("Exec Summary", "Snowflake Cortex", "One CDO-style headline paragraph"),
    ("Remediation", "SQL + Cortex", "Fix SQL + confidence + risk + approval flag; AI narrates"),
    ("Validation", "Deterministic SQL", "Re-detect + re-score, prove before \u2192 after"),
    ("Governance Recorder", "SQL", "Immutable audit trail of every action + approval"),
]
tbl = s.shapes.add_table(len(agents) + 1, 4, Inches(0.6), Inches(1.5), Inches(12.13), Inches(4.6)).table
for j, w in enumerate([0.55, 2.6, 2.4, 6.58]):
    tbl.columns[j].width = Inches(w)
for j, h in enumerate(["#", "Agent", "Tech", "Responsibility"]):
    cell(s, tbl, 0, j, h, NAVY, WHITE, bold=True)
for i, (a, t, r) in enumerate(agents, start=1):
    ai = "Cortex" in t
    fill = SF_LIGHT if ai else (WHITE if i % 2 else BG)
    cell(s, tbl, i, 0, str(i), fill, NAVY, bold=True)
    cell(s, tbl, i, 1, a, fill, NAVY, bold=True)
    cell(s, tbl, i, 2, t, fill, SF_MID if ai else TEXT, bold=ai)
    cell(s, tbl, i, 3, r, fill, TEXT)
for i in range(len(agents) + 1):
    tbl.rows[i].height = Inches(0.46)
b = box(s, 0.6, 6.3, 12.13, 0.62, fill=NAVY, line=None)
text(s, 0, 0, 0, 0, [[("Design principle:  ", {"color": SF_BLUE, "bold": True}),
                      ("SQL detects and executes (auditable). AI only reasons, explains and narrates. "
                       "A human approves anything risky.", {"color": WHITE})]],
     size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
notes(s, "Here are the nine agents. The blue rows use Snowflake Cortex, with mistral-large2 through "
         "CORTEX.COMPLETE, for the work that needs language and reasoning: root cause, business impact and an "
         "executive summary. The Remediation Agent pairs deterministic fix SQL with Cortex narration, and "
         "attaches a confidence score, a risk level and an approval flag to every fix. Detection, scoring, "
         "validation and governance are deterministic SQL. That's the design principle on the banner: SQL "
         "detects and executes so everything is auditable. AI only reasons and explains, and a human approves "
         "anything risky.")

# =============================================================================
# 6. Volume-aware scoring
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Volume-aware scoring: severity and spread", "The intelligence upgrade", 6)
box(s, 0.6, 1.55, 6.7, 2.55, fill=NAVY, line=None)
text(s, 0.9, 1.7, 6.2, 0.4, "THE FORMULA", size=12, color=SF_BLUE, bold=True)
text(s, 0.9, 2.1, 6.3, 1.9, [
    [("penalty", {"color": SF_BLUE}), ("  = PRESENCE(sev)", {})],
    [("         + VOLUME_MULT(sev) \u00d7 pct_rows", {})],
    [("table", {"color": SF_BLUE}), ("    = GREATEST(0, 100 \u2212 \u03a3 penalty)", {})],
    [("overall", {"color": SF_BLUE}), ("  = AVG(table scores)", {})],
], size=17, color=WHITE, font=MONO, spacing=1.35)

sev = [("Severity", "PRESENCE", "VOLUME_MULT / 1% rows"), ("CRITICAL", "8", "1.5"), ("HIGH", "4", "1.0"),
       ("MEDIUM", "2", "0.7"), ("LOW", "1", "0.3")]
sev_col = {"CRITICAL": RED, "HIGH": AMBER_DARK, "MEDIUM": SF_MID, "LOW": GREEN}
t = s.shapes.add_table(5, 3, Inches(7.65), Inches(1.55), Inches(5.08), Inches(2.55)).table
for j, w in enumerate([1.7, 1.38, 2.0]):
    t.columns[j].width = Inches(w)
for i, row in enumerate(sev):
    for j, v in enumerate(row):
        if i == 0:
            cell(s, t, i, j, v, NAVY, WHITE, size=12, bold=True,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
        else:
            cell(s, t, i, j, v, WHITE if i % 2 else BG, sev_col[v] if j == 0 else TEXT, size=14,
                 bold=(j == 0), align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    t.rows[i].height = Inches(0.51)

box(s, 0.6, 4.4, 6.0, 2.45)
text(s, 0.85, 4.55, 5.6, 0.4, "Why it matters", size=18, color=NAVY, bold=True)
bullets(s, 0.85, 5.0, 5.6, 1.8, [
    "A large, mostly-clean table is correctly rewarded.",
    "A structural breach (orphan / duplicate key) still hurts, even at low volume.",
    "This is how a real CDO reasons about trust.",
], size=14, gap=6)
box(s, 6.85, 4.4, 5.88, 2.45, fill=SF_LIGHT, line=None)
text(s, 7.1, 4.55, 5.4, 0.4, "Worked example", size=18, color=SF_MID, bold=True)
text(s, 7.1, 5.0, 5.4, 1.8, [
    "A CRITICAL issue touching 2% of rows:",
    [("8 + 1.5 \u00d7 2 = 11-point penalty", {"font": MONO, "bold": True, "color": NAVY, "size": 17})],
    "A LOW issue touching 2% of rows:",
    [("1 + 0.3 \u00d7 2 = 1.6-point penalty", {"font": MONO, "bold": True, "color": NAVY, "size": 17})],
], size=14, color=TEXT, spacing=1.25)
notes(s, "A naive score subtracts a flat penalty per issue type, so a ten-row table and a million-row table "
         "look the same. GuardianAI's Health Scorer is volume-aware. Each issue costs a fixed presence penalty "
         "by severity, plus a volume multiplier times the percentage of rows affected. Table health is 100 "
         "minus the total penalty, floored at zero, and the overall score is the average across the five "
         "tables. In practice, a critical issue on 2% of rows costs 11 points, while a low issue on the same "
         "rows costs 1.6. Big, mostly-clean tables are rewarded, but a structural breach like an orphan key "
         "still hurts even at low volume.")

# =============================================================================
# 7. Cross-table integrity
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Cross-table integrity: the differentiator", "Beyond single-table checks", 7)
box(s, 0.6, 1.55, 6.55, 5.3, fill=WHITE)
text(s, 0.85, 1.7, 6, 0.4, "Checks impossible from one table alone", size=15, color=MUTED, bold=True)
for name, (x, y, w) in {"CUSTOMERS": (0.85, 2.55, 1.6), "ORDERS": (3.05, 2.55, 1.6),
                        "PAYMENTS": (5.25, 2.55, 1.6), "PRODUCTS": (1.6, 5.2, 1.7),
                        "INVENTORY": (4.3, 5.2, 1.7)}.items():
    c = box(s, x, y, w, 0.75, fill=NAVY, line=None, radius=0.15)
    text(s, 0, 0, 0, 0, name, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
arrow_line(s, 3.03, 2.92, 2.48, 2.92, color=RED)      # ORDERS -> CUSTOMERS
text(s, 1.55, 3.4, 2.4, 0.35, "ORPHAN_ORDER", size=11, color=RED, bold=True, align=PP_ALIGN.CENTER)
arrow_line(s, 5.23, 2.92, 4.68, 2.92, color=RED)      # PAYMENTS -> ORDERS
text(s, 3.5, 3.4, 2.6, 0.6, ["ORPHAN_PAYMENT", "AMOUNT_MISMATCH"], size=11, color=RED, bold=True,
     align=PP_ALIGN.CENTER)
arrow_line(s, 4.28, 5.57, 3.33, 5.57, color=RED)      # INVENTORY -> PRODUCTS
text(s, 2.5, 4.8, 2.6, 0.35, "ORPHAN_INVENTORY", size=11, color=RED, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.85, 6.35, 6.1, 0.4, "Arrows point from child to parent table", size=11, color=MUTED, italic=True)

for i, (t, d) in enumerate([
    ("ORPHAN_ORDER", "An order whose customer doesn't exist."),
    ("ORPHAN_PAYMENT", "A payment whose order doesn't exist."),
    ("PAYMENT_AMOUNT_MISMATCH", "Payment amount \u2260 the matching order's amount."),
    ("ORPHAN_INVENTORY", "Stock for a product that doesn't exist."),
]):
    y = 1.55 + i * 1.34
    box(s, 7.45, y, 5.28, 1.2)
    box(s, 7.45, y, 0.1, 1.2, fill=RED, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 7.75, y + 0.15, 4.9, 0.4, t, size=16, color=NAVY, bold=True, font=MONO)
    text(s, 7.75, y + 0.6, 4.9, 0.5, d, size=14, color=MUTED)
notes(s, "This is what I think sets GuardianAI apart. Most data quality checks look at one table at a time. "
         "GuardianAI also checks relationships across the five tables. Orphan orders have a customer that "
         "doesn't exist. Orphan payments point to an order that doesn't exist. Payment amount mismatches are "
         "payments that don't equal the order they belong to. Orphan inventory is stock for a product that "
         "doesn't exist. These are exactly the issues that silently break revenue reporting, and you can't "
         "see them from a single table.")

# =============================================================================
# 8. Event-driven orchestration
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Event-driven orchestration: two loops", "Streams + Task DAG", 8)


def flow(slide, y, label, sub, steps, color, human_first=False):
    box(slide, 0.6, y, 12.13, 2.0, fill=WHITE)
    text(slide, 0.85, y + 0.15, 6, 0.4, label, size=19, color=NAVY, bold=True)
    pill(slide, 9.3, y + 0.17, 3.2, 0.4, sub, fill=AMBER_LIGHT if human_first else SF_LIGHT,
         color=AMBER_DARK if human_first else SF_MID, size=12)
    n, gap = len(steps), 0.32
    w = (11.6 - gap * (n - 1)) / n
    for i, st in enumerate(steps):
        x = 0.85 + i * (w + gap)
        c = box(slide, x, y + 0.8, w, 0.9, fill=AMBER if (human_first and i == 0) else color, line=None, radius=0.15)
        text(slide, 0, 0, 0, 0, st, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, shape=c)
        if i < n - 1:
            arrow_line(slide, x + w + 0.03, y + 1.25, x + w + gap - 0.03, y + 1.25, color=MUTED, width=1.75)


flow(s, 1.5, "Loop A: Ingest \u2192 Detect", "Fully automatic",
     ["File lands", "Stream fires", "Detect", "Score", "Rebuild Plan", "Analyze", "Dashboard updates"], NAVY)
flow(s, 3.7, "Loop B: Human action \u2192 Re-clean", "Human-gated: never self-fires",
     ["User clicks Approve", "Remediate (stream-safe)", "Validate", "Govern", "Dashboard updates"], SF_MID,
     human_first=True)
for (k, v), x, w in zip([("6", "Streams"), ("7", "Tasks"), ("18", "Stored procedures"),
                         ("SYSTEM$STREAM_HAS_DATA", "gating: a cheap no-op when idle")],
                        [0.6, 2.75, 4.9, 7.6], [2.0, 2.0, 2.55, 5.13]):
    box(s, x, 5.95, w, 0.95, fill=SF_LIGHT, line=None)
    big = len(k) < 4
    text(s, x + 0.1, 6.0, w - 0.2, 0.5, k, size=24 if big else 15, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER, font=FONT if big else MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.1, 6.48, w - 0.2, 0.35, v, size=12, color=SF_MID, align=PP_ALIGN.CENTER)
notes(s, "Orchestration is event-driven and uses only native Snowflake: six Streams, seven Tasks and eighteen "
         "stored procedures. Loop A is fully automatic. When a file lands, a Stream captures the new rows and "
         "the Task DAG runs detect, score, rebuild plan and analyse, and the dashboard updates. Loop B is "
         "human-gated and never fires on its own. Only when a user clicks Approve does it remediate, validate "
         "and write to the governance log. Tasks are gated on SYSTEM$STREAM_HAS_DATA, so when nothing has "
         "changed they're a cheap no-op, which protects credits.")

# =============================================================================
# 9. Responsible AI
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Responsible AI, by design", "Safe autonomy", 9)
for i, (t, d) in enumerate([
    ("Human-in-the-loop", "Every high-risk fix needs explicit approval."),
    ("Quarantine, never delete", "Bad rows are preserved for review."),
    ("Deterministic fix SQL", "No AI-generated code executes blindly."),
    ("Stream-safe remediation", "Dedup uses TRUNCATE + INSERT, never drops tables."),
    ("Full governance log", "Every agent action is recorded with an actor."),
]):
    y = 1.5 + i * 1.08
    box(s, 0.6, y, 5.0, 0.95)
    badge(s, 0.78, y + 0.2, 0.55, "\u2713", fill=GREEN, size=18)
    text(s, 1.5, y + 0.08, 4.0, 0.4, t, size=16, color=NAVY, bold=True)
    text(s, 1.5, y + 0.47, 4.0, 0.45, d, size=13, color=MUTED)
picture(s, os.path.join(SHOT, "6. Governance Log.png"), 5.95, 1.5, w=6.78)
text(s, 5.95, 5.55, 6.78, 0.4, "Governance Log: an immutable audit trail of every agent action and approval",
     size=12, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
notes(s, "Autonomy is only useful if it's safe, so Responsible AI is built into the architecture, not added "
         "on. Every high-risk fix needs explicit human approval. Bad rows are quarantined, never deleted. The "
         "fix SQL is deterministic and reviewable, so no AI-generated code runs blindly. Remediation is "
         "stream-safe: deduplication uses truncate-and-insert rather than dropping tables. Every action, by "
         "an agent or a human, is written to the governance log on the right, with an actor and a timestamp.")

# =============================================================================
# 10. Results
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Results: every table moves from HIGH risk to LOW", "Impact", 10)
results = [("CUSTOMERS", 51, 100), ("INVENTORY", 55, 100), ("PAYMENTS", 64, 100), ("PRODUCTS", 68, 100),
           ("ORDERS", 72, 100), ("OVERALL", 62, 100)]
cd = CategoryChartData()
cd.categories = [r[0] for r in results]
cd.add_series("Before", [r[1] for r in results])
cd.add_series("After", [r[2] for r in results])
box(s, 0.6, 1.5, 8.3, 5.4, fill=WHITE)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.75), Inches(1.65), Inches(8.0),
                        Inches(5.1), cd).chart
ch.font.name = FONT
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(13)
ch.legend.font.color.rgb = TEXT
va = ch.value_axis
va.maximum_scale, va.minimum_scale, va.major_unit = 110, 0, 25
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = BORDER
va.tick_labels.font.size = Pt(11)
va.tick_labels.font.color.rgb = MUTED
va.format.line.fill.background()
ca = ch.category_axis
ca.tick_labels.font.size = Pt(12)
ca.tick_labels.font.color.rgb = NAVY
ca.tick_labels.font.bold = True
ca.format.line.color.rgb = BORDER
plot = ch.plots[0]
plot.gap_width = 70
plot.overlap = -10
plot.has_data_labels = True
plot.data_labels.font.size = Pt(12)
plot.data_labels.font.bold = True
plot.data_labels.font.color.rgb = NAVY
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for ser, col in zip(plot.series, [RED, SF_BLUE]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col

box(s, 9.15, 1.5, 3.58, 2.45, fill=NAVY, line=None)
text(s, 9.3, 1.65, 3.3, 0.4, "OVERALL TRUST SCORE", size=12, color=SF_BLUE, bold=True, align=PP_ALIGN.CENTER)
text(s, 9.2, 2.1, 3.5, 1.0, [[("62", {"color": RGBColor(0xFF, 0x8A, 0x8E)}), (" \u2192 ", {"color": WHITE, "size": 28}),
                              ("~100", {"color": SF_BLUE})]], size=44, bold=True, align=PP_ALIGN.CENTER)
text(s, 9.3, 3.15, 3.3, 0.6, "after one human approval", size=13, color=SOFT, align=PP_ALIGN.CENTER)
for i, (k, v) in enumerate([("5 / 5", "tables HIGH \u2192 LOW risk"), ("0", "rows deleted: deduped or quarantined"),
                            ("100%", "of actions in the audit log")]):
    y = 4.15 + i * 0.93
    box(s, 9.15, y, 3.58, 0.82)
    text(s, 9.25, y + 0.05, 1.25, 0.72, k, size=20, color=SF_MID, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.5, y + 0.05, 2.15, 0.72, v, size=12, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Here are the results on the full dataset: about 1,000 customers, 5,000 orders, 500 products, 4,500 "
         "payments and 500 inventory rows, with a known set of injected issues so the run is reproducible. "
         "Before remediation, scores ranged from 51 for CUSTOMERS to 72 for ORDERS, for an overall score of 62. "
         "After one approval, every table recovers to 100 and the overall score reaches about 100. All five "
         "tables move from high to low business risk. No row is deleted: each one is either deduplicated or "
         "quarantined, and every step is in the audit log.")

# =============================================================================
# 11. The app
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "The GuardianAI app: Streamlit in Snowflake", "Live demo surface", 11)
for i, (t, d) in enumerate([
    ("Data Trust Overview", "Live trust score, per-table risk, before \u2192 after."),
    ("Findings Explorer", "Every issue, filterable by severity and table."),
    ("AI Investigation", "Cortex root cause + business impact per issue."),
    ("Remediation Center", "Fix, confidence, risk and the approval gate."),
    ("Governance Log", "Full audit trail of agent and human actions."),
]):
    y = 1.5 + i * 1.08
    box(s, 0.6, y, 3.35, 0.95)
    badge(s, 0.75, y + 0.22, 0.5, str(i + 1))
    text(s, 1.38, y + 0.07, 2.5, 0.38, t, size=14, color=NAVY, bold=True)
    text(s, 1.38, y + 0.42, 2.5, 0.52, d, size=11, color=MUTED, spacing=1.0)
GW, GH = 4.3, 2.52
for i, (f, cap) in enumerate([("1. Data Trust Overview.png", "Data Trust Overview"),
                              ("3. Findings Explorer.png", "Findings Explorer"),
                              ("4. AI Investigation Center.png", "AI Investigation Center"),
                              ("5. Remediation Center.png", "Remediation Center")]):
    x = 4.2 + (i % 2) * (GW + 0.23)
    y = 1.5 + (i // 2) * (GH + 0.38)
    picture(s, os.path.join(SHOT, f), x, y, w=GW, h=GH)
    text(s, x, y + GH + 0.02, GW, 0.3, cap, size=11, color=SF_MID, bold=True, align=PP_ALIGN.CENTER)
notes(s, "The whole system is visible through a five-page Streamlit-in-Snowflake app. The Overview shows the "
         "live trust score and per-table risk, including the before-and-after view. Findings Explorer lists "
         "every issue, filterable by severity and table. AI Investigation shows the Cortex-generated root cause "
         "and business impact for each issue. Remediation Center is where the proposed fix, its confidence and "
         "its risk are shown, and where a human approves it. The Governance Log, which we saw earlier, is the "
         "audit trail. One click on Approve all and run remediation, and the overall score climbs to about 100.")

# =============================================================================
# 12. Built on Snowflake + roadmap + close
# =============================================================================
s = prs.slides.add_slide(BLANK)
header(s, "Built on Snowflake, built to grow", "Stack, roadmap & thanks", 12)
box(s, 0.6, 1.5, 6.0, 4.15)
text(s, 0.85, 1.65, 5.5, 0.4, "Snowflake features used", size=18, color=NAVY, bold=True)
for i, (k, v) in enumerate([
    ("Snowflake Cortex", "CORTEX.COMPLETE (mistral-large2)"),
    ("Streams", "change data capture on new rows"),
    ("Tasks / DAG", "event-driven Loop A + Loop B"),
    ("Stored procedures", "agents as callable SQL units"),
    ("Streamlit in Snowflake", "5-page data trust app"),
    ("Stages + Snowsight", "ingest and landing"),
]):
    y = 2.15 + i * 0.57
    pill(s, 0.85, y, 2.45, 0.44, k, size=12)
    text(s, 3.4, y, 3.1, 0.44, v, size=13, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)

box(s, 6.85, 1.5, 5.88, 4.15)
text(s, 7.1, 1.65, 5.4, 0.4, "Roadmap", size=18, color=NAVY, bold=True)
for i, (t, d) in enumerate([
    ("Scheduled autonomous scans", "with Slack / Teams alerts on new criticals"),
    ("Learned severity weighting", "tuned from historical approvals"),
    ("Data contracts", "+ upstream prevention, not just repair"),
]):
    y = 2.2 + i * 1.12
    badge(s, 7.1, y + 0.08, 0.5, str(i + 1), fill=SF_BLUE)
    text(s, 7.75, y, 4.8, 0.4, t, size=16, color=NAVY, bold=True)
    text(s, 7.75, y + 0.4, 4.8, 0.45, d, size=13, color=MUTED)

box(s, 0.6, 5.9, 12.13, 1.0, fill=NAVY, line=None)
text(s, 0.9, 5.98, 8.0, 0.85, [[("Thank you", {"size": 26, "bold": True, "color": WHITE})],
                                [("GuardianAI turns messy data into trusted data, safely.",
                                  {"size": 14, "color": SOFT})]],
     anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
text(s, 8.9, 5.98, 3.6, 0.85, [[("Biswajit Jena", {"size": 18, "bold": True, "color": SF_BLUE})],
                               [("CoCo Quest 2026  |  Theme 1", {"size": 12, "color": SOFT})]],
     align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
notes(s, "GuardianAI is built entirely on native Snowflake. Cortex handles reasoning, Streams and Tasks handle "
         "event-driven orchestration, stored procedures implement the agents, and Streamlit in Snowflake is "
         "the front end. Nothing leaves the platform. Next on the roadmap: scheduled autonomous scans with "
         "Slack or Teams alerts on new critical issues, severity weights learned from historical approvals, "
         "and data contracts that prevent bad data upstream instead of only repairing it. Thank you. I'm "
         "Biswajit Jena, and I'm happy to take questions or walk through a live demo.")

# Save to local temp first (the workspace is a stage mount), then copy into place.
tmp = os.path.join(tempfile.gettempdir(), "coco_quest.pptx")
prs.save(tmp)
shutil.copyfile(tmp, OUT)
print("saved", OUT, len(prs.slides), "slides")
